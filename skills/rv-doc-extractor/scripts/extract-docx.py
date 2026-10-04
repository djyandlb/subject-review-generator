#!/usr/bin/env python
"""DOCX 提取（python-docx）：段落 + 表格（表格单元格以 | 连接）
用法: python extract-docx.py <源.docx> <输出.txt>
依赖: pip install python-docx
"""
import sys

from docx import Document


def main():
    if len(sys.argv) < 3:
        print("用法: python extract-docx.py <源.docx> <输出.txt>")
        sys.exit(1)
    src, out = sys.argv[1], sys.argv[2]
    doc = Document(src)
    parts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    print("DOCX → %s" % out)


if __name__ == "__main__":
    main()

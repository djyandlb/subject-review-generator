#!/usr/bin/env python
"""PPTX 提取（python-pptx）：每张幻灯片文本框 + 备注
用法: python extract-pptx.py <源.pptx> <输出.txt>
依赖: pip install python-pptx
"""
import sys

from pptx import Presentation


def main():
    if len(sys.argv) < 3:
        print("用法: python extract-pptx.py <源.pptx> <输出.txt>")
        sys.exit(1)
    src, out = sys.argv[1], sys.argv[2]
    prs = Presentation(src)
    parts = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        texts.append(t)
        # 演讲者备注（notes）：资料关键信息可能写在备注里，一并提取（个别 PPT 备注异常不阻断）
        try:
            if slide.has_notes_slide:
                nf = slide.notes_slide.notes_text_frame
                nt = "\n".join(p.text.strip() for p in nf.paragraphs if p.text.strip())
                if nt:
                    texts.append("【备注】\n" + nt)
        except Exception:
            pass
        parts.append("=== SLIDE %d ===\n%s" % (i + 1, "\n".join(texts)))
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    print("PPTX %d 页 → %s" % (len(prs.slides), out))


if __name__ == "__main__":
    main()
